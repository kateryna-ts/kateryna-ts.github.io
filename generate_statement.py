"""
Generate a Statement of Career Goals Word document (.docx) for
Kateryna Tsekhmayster's PhD in Finance application at FAU.
Run: python3 generate_statement.py
Output: Kateryna_Tsekhmayster_Statement_of_Career_Goals.docx
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAVY  = RGBColor(0x1B, 0x2A, 0x4A)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY  = RGBColor(0x55, 0x55, 0x55)
FONT  = "Calibri"


def styled(run, size=12, bold=False, italic=False, color=BLACK):
    run.font.name      = FONT
    run.font.size      = Pt(size)
    run.font.bold      = bold
    run.font.italic    = italic
    run.font.color.rgb = color


def body_para(doc, text, space_before=0, space_after=10, first_line=0.3):
    p = doc.add_paragraph()
    p.paragraph_format.space_before       = Pt(space_before)
    p.paragraph_format.space_after        = Pt(space_after)
    p.paragraph_format.first_line_indent  = Inches(first_line)
    p.paragraph_format.line_spacing       = Pt(14)
    r = p.add_run(text)
    styled(r, 12, color=BLACK)
    return p


def build_statement():
    doc = Document()

    for section in doc.sections:
        section.top_margin    = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin   = Inches(1.25)
        section.right_margin  = Inches(1.25)

    # ── Title block ────────────────────────────────────────────────────────────
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after  = Pt(4)
    title_r = title_p.add_run("Statement of Career Goals")
    styled(title_r, 16, bold=True, color=NAVY)

    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_p.paragraph_format.space_before = Pt(0)
    name_p.paragraph_format.space_after  = Pt(2)
    name_r = name_p.add_run("Kateryna Tsekhmayster")
    styled(name_r, 12, bold=True, color=BLACK)

    prog_p = doc.add_paragraph()
    prog_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prog_p.paragraph_format.space_before = Pt(0)
    prog_p.paragraph_format.space_after  = Pt(20)
    prog_r = prog_p.add_run("PhD in Finance — Florida Atlantic University")
    styled(prog_r, 11, italic=True, color=GRAY)

    # ── Body ───────────────────────────────────────────────────────────────────
    paragraphs = [

        # 1 — Opening hook
        (
            "Quantum mechanics was once thought to belong exclusively to the realm of physics. Today, it is "
            "poised to reshape the foundations of financial economics. The emergence of quantum computing "
            "introduces capabilities—superposition, entanglement, and quantum parallelism—that can solve "
            "optimization and simulation problems that remain intractable for classical machines. My "
            "academic and research trajectory has converged on a single conviction: the next generation of "
            "financial economists must be equipped to harness these capabilities, and the institutions "
            "training them must build bridges between quantum science and empirical finance. I am applying "
            "to the PhD in Finance program at Florida Atlantic University to become that kind of scholar."
        ),

        # 2 — Quantum research genesis
        (
            "My engagement with quantum economics grew out of a broader inquiry into why some countries "
            "fall behind in transformative technological transitions. While pursuing my dual Master's "
            "degrees in Economics and Political Science at FAU, I began examining how workforce readiness "
            "determines a nation's capacity to absorb and deploy emerging technologies. This line of "
            "inquiry culminated in my paper, \"Quantum Labor Market Development in Emerging Countries,\" "
            "presented at FAU Graduate Research Day. In that work I developed the Quantum Labor Market "
            "Readiness Index (QLRI), a composite framework that scores nations across four dimensions: "
            "digital infrastructure, human capital, innovation ecosystem, and governance quality. The QLRI "
            "quantifies readiness gaps that standard productivity metrics overlook and provides a "
            "diagnostic tool for policymakers seeking to position their economies for quantum-era "
            "competition. Presenting this research at FAU Graduate Research Day sharpened my ambition to "
            "extend the framework into financial market applications—specifically, how quantum readiness "
            "differentials translate into sovereign risk premiums, capital flow asymmetries, and "
            "cross-border investment patterns."
        ),

        # 3 — Broader research portfolio
        (
            "My research portfolio reflects a consistent commitment to applying rigorous analytical methods "
            "to high-stakes problems at the intersection of economics, finance, and geopolitics. In "
            "\"Freedom as Development: Economics of the Belarusian Brain Drain and Remittances\" "
            "(SSRN: 10.2139/ssrn.5467415), I used Amartya Sen's capabilities framework to model skilled "
            "emigration as an endogenous response to political repression, examining the macroeconomic "
            "feedback loops between institutional quality and human capital flight. In \"Brazilian Oil, "
            "Petrobras, and US-China Rivalry,\" I analyzed how a state-owned enterprise's strategic "
            "positioning reshapes bilateral financial relationships and commodity pricing dynamics amid "
            "great-power competition. Each project has reinforced my sense that financial decisions are "
            "embedded in institutional and geopolitical structures that purely market-centered models "
            "routinely underspecify."
        ),

        # 4 — Why Finance PhD, why FAU
        (
            "I am drawn to a PhD in Finance—rather than Economics broadly—because I want to develop "
            "technical depth in asset pricing, market microstructure, and quantitative modeling that will "
            "make my quantum economics work genuinely tractable. The Finance PhD curriculum's emphasis on "
            "stochastic processes, portfolio theory, and empirical methods provides exactly the toolkit "
            "I need to translate QLRI-derived readiness scores into testable hypotheses about capital "
            "markets. Florida Atlantic University is my institution of choice for this transition. FAU's "
            "College of Business is home to faculty whose work on financial markets, international "
            "finance, and financial technology aligns directly with my research agenda. I have already "
            "experienced FAU's intellectual environment through my Master's programs and Graduate Research "
            "Day, and I am eager to deepen my engagement with the university's community of scholars. "
            "The proximity to South Florida's growing fintech and international banking sector further "
            "enriches the applied dimension of my planned research."
        ),

        # 5 — Dissertation vision
        (
            "My proposed doctoral research agenda centers on three interconnected questions. First, do "
            "quantum readiness differentials—as measured by QLRI or analogous indices—predict sovereign "
            "bond spreads and equity risk premiums across emerging markets? Second, can quantum-inspired "
            "optimization algorithms improve portfolio construction in high-dimensional, illiquid markets "
            "where classical mean-variance approaches break down? Third, how do technological leadership "
            "gaps in quantum computing interact with existing financial integration patterns to produce "
            "new forms of capital-market stratification between advanced and developing economies? These "
            "questions sit at the frontier of financial economics and have direct implications for "
            "international investors, central banks, and development finance institutions navigating the "
            "quantum transition."
        ),

        # 6 — Interdisciplinary strengths
        (
            "My interdisciplinary training positions me to pursue this agenda from multiple angles. My "
            "work on blockchain-based conflict resolution (\"Blockchain of International Relations: "
            "Universal Cosmopolitan Protocol,\" SSRN: 10.2139/ssrn.5272407) gave me hands-on experience "
            "modeling distributed ledger architectures—a skill transferable to decentralized finance "
            "applications and the tokenization of quantum-secured assets. My political science background "
            "equips me to contextualize financial market outcomes within governance and institutional "
            "frameworks, an increasingly valuable perspective as regulators worldwide grapple with "
            "quantum cryptography standards and their implications for financial infrastructure. My "
            "language skills in English, Ukrainian, Russian, and French expand my ability to conduct "
            "primary source research across post-Soviet and Francophone emerging markets that are "
            "underrepresented in the quantum finance literature."
        ),

        # 7 — Career trajectory
        (
            "Upon completing the PhD, my primary goal is an academic career at a research university "
            "where I can build a laboratory focused on quantum financial economics and mentor the next "
            "cohort of scholars entering this field. I envision publishing in top-tier finance and "
            "economics journals and contributing to the emerging policy literature on quantum readiness "
            "and financial stability. In parallel, I intend to engage with international institutions—"
            "the IMF, World Bank, and regional development banks—that are beginning to assess the "
            "financial stability implications of uneven quantum adoption across their member economies. "
            "A secondary pathway I find compelling is applied research in the growing quantum fintech "
            "sector, where the gap between theoretical quantum algorithms and real-world portfolio "
            "management remains wide and consequential."
        ),

        # 8 — Closing
        (
            "The convergence of quantum computing, international finance, and development economics "
            "defines a research frontier that is both intellectually exciting and practically urgent. "
            "I have spent the past several years building the interdisciplinary foundations to work "
            "at that frontier, and I am ready to commit to the rigorous doctoral training that will "
            "make that work credible and impactful. I am grateful for the opportunity to apply to "
            "Florida Atlantic University's PhD in Finance program, and I look forward to contributing "
            "to its scholarly community."
        ),
    ]

    for text in paragraphs:
        body_para(doc, text, space_after=10)

    # ── Save ───────────────────────────────────────────────────────────────────
    output = "Kateryna_Tsekhmayster_Statement_of_Career_Goals.docx"
    doc.save(output)
    print(f"Saved: {output}")


if __name__ == "__main__":
    build_statement()
